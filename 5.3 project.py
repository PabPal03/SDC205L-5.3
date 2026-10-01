from datetime import datetime
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference

def convertData(fahrenheit):
    return (fahrenheit - 32) * 5 / 9

def insertData(filePath, data):
    try:
        file = open(filePath, "a")
        file.write(data + "\n")
        file.close()
    except:
        print("Error writing to the file")

def viewData(filePath):
    print(f"The file {filePath}")
    try:
        file = open(filePath, "r")
        print(file.read())
        file.close()
    except:
        print("Error reading the file")

def getInput():
    entries = int(input("How many entries are you inputting? "))
    for i in range(entries):
        entryDate = input("Enter a date: ")
        value = float(input("Enter the highest temp for the inputted date: "))
        celsius = convertData(value)
        data = f"{entryDate},{value},{celsius:.1f}"
        try:
            insertData("ZooData.csv", data)
            print(f"The following data was saved at {datetime.now()}: {data}.")
        except:
            print("Error saving data")

def createChart(filePath, chartType):
    source = input("Choose data source (Fahrenheit or Celsius): ")
    dates = []
    values = []
    try:
        file = open(filePath, "r")
        for line in file:
            parts = line.strip().split(",")
            if len(parts) < 3:
                continue
            dates.append(parts[0])
            if source.lower() == "celsius":
                values.append(float(parts[2]))
            else:
                values.append(float(parts[1]))
        file.close()
    except:
        print("Error reading the file")
        return

    wb = Workbook()
    sheet = wb.active
    sheet["A1"] = "Date"
    sheet["B1"] = source
    for i in range(len(dates)):
        sheet.cell(row=i + 2, column=1, value=dates[i])
        sheet.cell(row=i + 2, column=2, value=values[i])

    if chartType.lower() == "bar":
        chart = BarChart()
    else:
        chart = LineChart()
    data = Reference(sheet, min_col=2, min_row=1, max_row=len(dates) + 1)
    labels = Reference(sheet, min_col=1, min_row=2, max_row=len(dates) + 1)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(labels)
    chart.x_axis.title = "Date"
    chart.y_axis.title = source
    chart.title = f"Pabpal1418 {datetime.now().strftime('%m/%d/%Y')}"
    sheet.add_chart(chart, "D2")
    try:
        wb.save("final.xlsx")
        print("Report saved to final.xlsx")
    except:
        print("Error saving final.xlsx. Close the file if it is open, then try again.")

def generateReport(filePath):
    chartType = input("Choose a graph type (line or bar): ")
    createChart(filePath, chartType)

while True:
    print("Pabpal1418's Spreadsheet Automation Menu")
    print("Choose a number from the following options")
    print("1 Input Data")
    print("2 View Current Data")
    print("3 Generate Report")
    print("4 Exit")

    choice = input("Enter your Choice: ")

    if choice == "1":
        print(f"You selected {choice} at {datetime.now()}")
        getInput()
    elif choice == "2":
        print(f"You selected {choice} at {datetime.now()}")
        viewData("ZooData.csv")
    elif choice == "3":
        print(f"You selected {choice} at {datetime.now()}")
        generateReport("ZooData.csv")
    elif choice == "4":
        break
    else:
        print("Error: The chosen functionality is not implemented yet")
